import hashlib
import logging
from pathlib import Path

import docx
from pdfminer.high_level import extract_text
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)


class IngestionError(Exception):
    """Base class for document ingestion failures."""


class UnsupportedFormatError(IngestionError):
    """Raised for files whose extension is not a supported document format."""


class DocumentLoadError(IngestionError):
    """Raised when a supported file cannot be parsed."""


class NoDocumentsFoundError(IngestionError):
    """Raised when the input directory yields no documents."""


class Document(BaseModel):
    """A single ingested source document."""

    path: Path
    text: str
    fingerprint: str = Field(description="SHA-256 of the extracted text")


def compute_fingerprint(text: str) -> str:
    """Fingerprint the extracted text, independent of the source format."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _load_plain_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def _load_pdf(path: Path) -> str:
    return extract_text(str(path))


def _load_docx(path: Path) -> str:
    paragraphs = docx.Document(str(path)).paragraphs
    return "\n".join(paragraph.text for paragraph in paragraphs)


LOADERS = {
    ".txt": _load_plain_text,
    ".md": _load_plain_text,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
}


def load_document(path: Path) -> Document:
    """Load a single supported document and fingerprint its text."""
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise UnsupportedFormatError(f"unsupported file format: {path.name}")
    try:
        text = loader(path)
    except Exception as error:
        raise DocumentLoadError(f"failed to load {path}: {error}") from error
    return Document(path=path, text=text, fingerprint=compute_fingerprint(text))


def load_documents(input_dir: Path) -> list[Document]:
    """Load every supported document under input_dir, recursively.

    Unreadable and unsupported files are logged and skipped so that a single
    bad file cannot abort ingestion.
    """
    if not input_dir.is_dir():
        raise NoDocumentsFoundError(f"input directory does not exist: {input_dir}")

    documents: list[Document] = []
    for path in sorted(input_dir.rglob("*")):
        if not path.is_file():
            continue
        try:
            documents.append(load_document(path))
        except UnsupportedFormatError as error:
            logger.warning("Skipping %s: %s", path, error)
        except DocumentLoadError as error:
            logger.error("Skipping %s: %s", path, error)

    if not documents:
        raise NoDocumentsFoundError(f"no supported documents found in {input_dir}")
    return documents
